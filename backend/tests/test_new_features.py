import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document as WordDocument
from fastapi.testclient import TestClient
from langchain_core.documents import Document as LCDocument
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, ChatMessage, Document, DocumentChunk, get_db
from app.main import app, process_document
from app.rag import index_document, search_documents


class NewFeaturesTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(bind=self.engine)

        def test_db():
            with self.session_factory() as session:
                yield session

        app.dependency_overrides[get_db] = test_db
        self.client = TestClient(app)
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        self.engine.dispose()
        self.temp_dir.cleanup()

    def test_general_chat_creates_conversation_without_sources(self):
        with patch("app.main.answer_general", return_value="Java is a programming language.") as answer:
            response = self.client.post("/api/chat", json={"question": "What is Java?", "mode": "general"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["sources"], [])
        answer.assert_called_once_with("What is Java?", history="")
        chat_id = response.json()["conversation_id"]
        history = self.client.get(f"/api/history?conversation_id={chat_id}").json()
        self.assertEqual([row["role"] for row in history], ["user", "assistant"])
        self.assertEqual(history[1]["sources"], [])
        self.assertEqual(self.client.get("/api/conversations").json()[0]["mode"], "general")

        with patch("app.main.answer_general", return_value="It runs on the JVM.") as follow_up:
            second = self.client.post("/api/chat", json={"question": "Where does it run?", "conversation_id": chat_id})
        self.assertEqual(second.status_code, 200)
        self.assertIn("user: What is Java?", follow_up.call_args.kwargs["history"])

    def test_document_citations_survive_history_reload(self):
        with self.session_factory() as session:
            session.add(Document(filename="notes.pdf", stored_path="notes.pdf", chunk_count=1))
            session.commit()
        source = LCDocument(page_content="Source passage", metadata={"document_id": 1, "filename": "notes.pdf", "page": 2})
        with patch("app.main.answer_question", return_value=("Answer [notes.pdf p2]", [source])):
            response = self.client.post("/api/chat", json={"question": "Question", "document_id": 1})
        self.assertEqual(response.status_code, 200)
        history = self.client.get(f"/api/history?conversation_id={response.json()['conversation_id']}").json()
        self.assertEqual(history[1]["sources"][0]["page"], 2)
        self.assertEqual(history[1]["sources"][0]["snippet"], "Source passage")

    def test_upload_validation_and_pending_status(self):
        path = Path(self.temp_dir.name) / "uploaded.txt"
        with patch("app.main.unique_upload_path", return_value=path), patch("app.main.process_document") as process:
            response = self.client.post("/api/documents", files={"file": ("notes.txt", b"Hello document", "text/plain")})
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.json()["status"], "pending")
        self.assertEqual(path.read_text(), "Hello document")
        process.assert_called_once()
        bad = self.client.post("/api/documents", files={"file": ("bad.exe", b"invalid", "application/octet-stream")})
        self.assertEqual(bad.status_code, 400)
        with patch("app.main.unique_upload_path", return_value=Path(self.temp_dir.name) / "bad.pdf"):
            bad_pdf = self.client.post("/api/documents", files={"file": ("bad.pdf", b"invalid", "application/pdf")})
        self.assertEqual(bad_pdf.status_code, 400)

    def test_selected_document_filter_applies_inside_search(self):
        from app.rag import _norm, _pack_embedding

        with self.session_factory() as session:
            session.add_all([
                Document(id=7, filename="selected.txt", stored_path="selected.txt"),
                Document(id=8, filename="other.txt", stored_path="other.txt"),
            ])
            session.add_all([
                DocumentChunk(document_id=7, filename="selected.txt", content="Selected", embedding=_pack_embedding([1.0, 0.0]), embedding_norm=_norm([1.0, 0.0])),
                DocumentChunk(document_id=8, filename="other.txt", content="Other", embedding=_pack_embedding([1.0, 0.0]), embedding_norm=_norm([1.0, 0.0])),
            ])
            session.commit()
            embeddings = MagicMock()
            embeddings.embed_query.return_value = [1.0, 0.0]
            with patch("app.rag.get_embeddings", return_value=embeddings):
                results = search_documents("Question", db=session, document_id=7)
        self.assertEqual([result.page_content for result in results], ["Selected"])

    def test_word_and_text_extraction(self):
        word_path = Path(self.temp_dir.name) / "note.docx"
        word = WordDocument()
        word.add_paragraph("A Word paragraph")
        word.save(word_path)
        text_path = Path(self.temp_dir.name) / "note.txt"
        text_path.write_text("A text paragraph", encoding="utf-8")
        for path, expected in [(word_path, "A Word paragraph"), (text_path, "A text paragraph")]:
            with self.subTest(path=path.suffix), self.session_factory() as session:
                session.add(Document(id=5, filename=path.name, stored_path=str(path)))
                session.commit()
                embeddings = MagicMock()
                embeddings.embed_documents.return_value = [[1.0, 0.0]]
                with patch("app.rag.get_embeddings", return_value=embeddings):
                    self.assertEqual(index_document(path, 5, path.name, session), 1)
                chunk = session.query(DocumentChunk).one()
                self.assertIn(expected, chunk.content)
                self.assertEqual(chunk.document_id, 5)
                self.assertIsNone(chunk.page)
                session.query(DocumentChunk).delete()
                session.query(Document).delete()
                session.commit()

    def test_background_index_failure_is_visible_and_retryable(self):
        path = Path(self.temp_dir.name) / "scan.pdf"
        path.write_bytes(b"%PDF-1.4")
        with self.session_factory() as session:
            session.add(Document(filename="scan.pdf", stored_path=str(path), status="pending"))
            session.commit()
        with patch("app.main.SessionLocal", self.session_factory), patch("app.main.index_document", side_effect=ValueError("No readable text found")):
            import asyncio
            asyncio.run(process_document(1))
        row = self.client.get("/api/documents").json()[0]
        self.assertEqual(row["status"], "failed")
        self.assertIn("OCR", row["error"])
        with patch("app.main.process_document") as run:
            response = self.client.post("/api/documents/1/retry")
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.json()["status"], "pending")
        run.assert_called_once_with(1)

    def test_legacy_chat_migration(self):
        legacy = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        try:
            with legacy.begin() as connection:
                connection.execute(text("CREATE TABLE documents (id INTEGER PRIMARY KEY, filename VARCHAR(255), stored_path TEXT, chunk_count INTEGER, created_at DATETIME)"))
                connection.execute(text("CREATE TABLE chat_messages (id INTEGER PRIMARY KEY, document_id INTEGER, role VARCHAR(20), content TEXT, created_at DATETIME)"))
                connection.execute(text("INSERT INTO chat_messages (role, content) VALUES ('user', 'Old question')"))
            with patch("app.database.engine", legacy):
                from app.database import init_db
                init_db()
            with legacy.connect() as connection:
                self.assertEqual(connection.execute(text("SELECT status FROM documents LIMIT 1")).fetchall(), [])
                self.assertIsNotNone(connection.execute(text("SELECT conversation_id FROM chat_messages")).scalar())
                self.assertEqual(connection.execute(text("SELECT title FROM conversations")).scalar(), "Previous chat")
        finally:
            legacy.dispose()


if __name__ == "__main__":
    unittest.main()
